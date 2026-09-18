# -*- coding: utf-8 -*-
"""
共用邏輯：照片 → 竣工查驗照片／驗收相片 Excel。
檔名開頭底線，Streamlit 不會把它當成一個頁面。
"""
import io
import os
import re
import shutil
import tempfile

import streamlit as st
from _auth import require_password
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Side, Font
from PIL import Image, ImageOps

# ---------------------------------------------------------------- 版面設定（依實際範本量測）
PRESETS = {
    "竣工查驗照片": dict(
        sheet_name="竣工查驗照片", title="竣工查驗相片",
        title_h=34.15, proj_h=25.15, caption_h=25.15, photo_h=220.15,
        col_w=45.75, caption_position="above",
    ),
    "驗收相片": dict(
        sheet_name="驗收相片", title="竣工驗收相片",
        title_h=27.75, proj_h=25.15, caption_h=24.95, photo_h=210,
        col_w=46.5, caption_position="above",
    ),
}

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(name="標楷體", size=18, bold=True)
PROJECT_FONT = Font(name="標楷體", size=12)
CAPTION_FONT = Font(name="標楷體", size=12)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _to_png(blob):
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
        im = ImageOps.exif_transpose(im)  # 依相機的 EXIF 方向轉正，保留原本橫直式
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGB")
        out = io.BytesIO()
        im.save(out, format="PNG")
        return out.getvalue(), im.size
    except Exception:
        return None, None


def _excel_col_width_px(chars, mdw=7):
    return int((256 * chars + int(128 / mdw)) / 256 * mdw)


def _fit_image(png_bytes, max_w_px, max_h_px):
    im = Image.open(io.BytesIO(png_bytes))
    w, h = im.size
    scale = min(max_w_px / w, max_h_px / h, 1.0)
    im = im.resize((max(1, int(w * scale)), max(1, int(h * scale))))
    out = io.BytesIO()
    im.save(out, format="PNG")
    out.seek(0)
    return out


def build_gallery_excel(photos, project_name, preset_key):
    preset = PRESETS[preset_key]
    wb = Workbook()
    # 對齊範本的預設字型（新細明體 12），欄寬字元數才會換算成一樣的像素寬度
    wb._named_styles[0].font = Font(name="新細明體", size=12)
    ws = wb.active
    ws.title = preset["sheet_name"]
    ws.column_dimensions["A"].width = preset["col_w"]
    ws.column_dimensions["B"].width = preset["col_w"]

    max_w_px = _excel_col_width_px(preset["col_w"]) - 8
    max_h_px = int(preset["photo_h"] * 4 / 3) - 8

    row = 1
    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"] = preset["title"]
    ws[f"A{row}"].font = TITLE_FONT
    ws[f"A{row}"].alignment = CENTER
    ws.row_dimensions[row].height = preset["title_h"]
    row += 1

    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"] = f"工程名稱：{project_name}"
    ws[f"A{row}"].font = PROJECT_FONT
    ws[f"A{row}"].alignment = LEFT
    ws.row_dimensions[row].height = preset["proj_h"]
    row += 1

    # photos 是「每一格」的清單（沒上傳照片的格子是 None），依原始格子編號兩個一列配對，
    # 不會因為中間有空格就把後面的照片往前遞補、跟原本的說明錯位
    count = 0
    for i in range(0, len(photos), 2):
        pair = photos[i:i + 2]
        if all(p is None for p in pair):
            continue

        def write_photo_row(r):
            ws.row_dimensions[r].height = preset["photo_h"]
            for col_letter, p in zip(("A", "B"), pair):
                cell = ws[f"{col_letter}{r}"]
                cell.border = BORDER
                if p is not None:
                    img = XLImage(_fit_image(p["png"], max_w_px, max_h_px))
                    img.anchor = f"{col_letter}{r}"
                    ws.add_image(img)

        def write_caption_row(r):
            ws.row_dimensions[r].height = preset["caption_h"]
            for idx, col_letter in enumerate(("A", "B")):
                cell = ws[f"{col_letter}{r}"]
                cell.border = BORDER
                cell.font = CAPTION_FONT
                cell.alignment = CENTER
                p = pair[idx] if idx < len(pair) else None
                cell.value = p["caption"] if p is not None else ""

        if preset["caption_position"] == "above":
            write_caption_row(row)
            row += 1
            write_photo_row(row)
        else:
            write_photo_row(row)
            row += 1
            write_caption_row(row)
        count += sum(1 for p in pair if p is not None)
        ws.cell(row=row, column=3, value=count)
        row += 1

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def merge_sheet_into_bytes(photos, project_name, preset_key, target_bytes):
    """用 Excel COM 把新產生的分頁取代掉 target_bytes 裡的同名分頁；其餘分頁與順序不動。
    回傳合併後的檔案 bytes，原始上傳內容完全不會被修改。"""
    import pythoncom
    import win32com.client as win32

    sheet_name = PRESETS[preset_key]["sheet_name"]
    tmp_dir = tempfile.mkdtemp(prefix="gallery2excel_")
    try:
        src_path = os.path.join(tmp_dir, "source.xlsx")
        with open(src_path, "wb") as f:
            f.write(build_gallery_excel(photos, project_name, preset_key).read())

        tgt_path = os.path.join(tmp_dir, "target.xlsx")
        with open(tgt_path, "wb") as f:
            f.write(target_bytes)

        pythoncom.CoInitialize()
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        try:
            src_wb = excel.Workbooks.Open(src_path)
            tgt_wb = excel.Workbooks.Open(tgt_path)
            try:
                idx = None
                for i in range(1, tgt_wb.Sheets.Count + 1):
                    if tgt_wb.Sheets(i).Name == sheet_name:
                        idx = i
                        break
                if idx is not None:
                    tgt_wb.Sheets(idx).Delete()

                src_sheet = src_wb.Sheets(1)
                if idx is not None and idx <= tgt_wb.Sheets.Count:
                    src_sheet.Copy(Before=tgt_wb.Sheets(idx))
                    new_sheet = tgt_wb.Sheets(idx)
                else:
                    src_sheet.Copy(After=tgt_wb.Sheets(tgt_wb.Sheets.Count))
                    new_sheet = tgt_wb.Sheets(tgt_wb.Sheets.Count)
                new_sheet.Name = sheet_name

                tgt_wb.Save()
            finally:
                tgt_wb.Close(False)
                src_wb.Close(False)
        finally:
            excel.Quit()
            pythoncom.CoUninitialize()

        with open(tgt_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _guess_project_name(filename):
    """從照片檔名猜工程名稱：取副檔名前、第一個底線之前的文字，
    並去掉開頭的編號前綴（例如「31-」）"""
    if not filename:
        return ""
    stem = os.path.splitext(filename)[0]
    name = stem.split("_")[0].strip()
    name = re.sub(r"^\d+-", "", name)
    return name


def render_page(preset_key):
    """畫出整個頁面（固定格子上傳照片、產生 Excel、合併），preset_key 固定成呼叫端指定的格式。"""
    sheet_name = PRESETS[preset_key]["sheet_name"]
    ns = preset_key  # session_state 的 key 前綴，讓不同頁面的狀態互不干擾

    st.set_page_config(page_title=f"照片 → {sheet_name} Excel", layout="wide")
    require_password()
    st.title(f"照片 → {sheet_name} Excel")

    slot_count_key = f"{ns}_slot_count"
    if slot_count_key not in st.session_state:
        st.session_state[slot_count_key] = 6
    n = st.session_state[slot_count_key]

    # 第一次偵測到已上傳的照片時，自動用檔名猜工程名稱（之後仍可手動修改）
    name_key = f"{ns}_project_name"
    auto_flag_key = f"{ns}_project_name_auto_done"
    if not st.session_state.get(auto_flag_key):
        for i in range(n):
            f = st.session_state.get(f"{ns}_slot_file_{i}")
            if f is not None:
                guessed = _guess_project_name(f.name)
                if guessed:
                    st.session_state[name_key] = guessed
                st.session_state[auto_flag_key] = True
                break

    with st.sidebar:
        st.header("設定")
        project_name = st.text_input("工程名稱（自動抓第一張照片檔名，可修改）", key=name_key)

    caption_above = PRESETS[preset_key]["caption_position"] == "above"

    st.subheader("① 逐格上傳照片")
    slots = []
    display_cols = 3
    for start in range(0, n, display_cols):
        cols = st.columns(display_cols)
        for j, col in enumerate(cols):
            i = start + j
            if i >= n:
                break
            with col:
                st.caption(f"格 {i + 1}")

                def _upload_widget():
                    f = st.file_uploader("照片", type=["jpg", "jpeg", "png"],
                                         key=f"{ns}_slot_file_{i}", label_visibility="collapsed")
                    if f is not None:
                        st.image(f, use_container_width=True)
                    return f

                def _caption_widget():
                    return st.text_input("說明", key=f"{ns}_slot_cap_{i}", label_visibility="collapsed",
                                         placeholder="說明")

                if caption_above:
                    cap = _caption_widget()
                    f = _upload_widget()
                else:
                    f = _upload_widget()
                    cap = _caption_widget()
                slots.append((f, cap))

    if st.button("＋ 新增格子", key=f"{ns}_add_slot"):
        st.session_state[slot_count_key] += 1
        st.rerun()

    st.divider()
    confirmed_key = f"{ns}_confirmed"
    if st.button("✅ 完成，準備產生 Excel", key=f"{ns}_confirm_btn", type="primary"):
        st.session_state[confirmed_key] = True

    if not st.session_state.get(confirmed_key):
        st.info("填好照片與說明後，按上面「✅ 完成，準備產生 Excel」再繼續。")
        return

    target_file = st.file_uploader("② 上傳結算明細表.xlsx（要把照片合併進去的檔案，選填）",
                                   type=["xlsx"], key=f"{ns}_target_xlsx")

    # 保留每一格原本的順序（沒上傳照片的格子放 None），輸出時才不會因為中間有空格
    # 就把後面照片的說明錯位配對到別張照片上
    photos = []
    for i, (f, cap) in enumerate(slots):
        if f is None:
            photos.append(None)
            continue
        png, size = _to_png(f.getvalue())
        if png is None:
            st.error(f"格 {i + 1} 的照片無法讀取，已略過。")
            photos.append(None)
            continue
        photos.append({"png": png, "size": size, "caption": cap})

    filled_count = sum(1 for p in photos if p is not None)

    st.divider()
    st.subheader(f"目前有 {filled_count} 張照片")
    if filled_count:
        xlsx = build_gallery_excel(photos, project_name, preset_key)
        st.download_button(
            "產生並下載 Excel",
            data=xlsx,
            file_name=(project_name or sheet_name) + ".xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"{ns}_download",
        )

        st.divider()
        st.subheader("合併到既有結算明細表")
        if target_file is None:
            st.info("要合併的話，請在上面 ② 上傳結算明細表.xlsx。")
        else:
            st.caption(f"會用目前的 {filled_count} 張照片，取代掉「{target_file.name}」裡"
                       f"「{sheet_name}」那個分頁，其他分頁不動，完成後另外下載，"
                       "不會動到你上傳的原始檔案。需要這台電腦有安裝 Microsoft Excel。")
            if st.button("合併並產生新檔案", type="primary", key=f"{ns}_merge_btn"):
                with st.spinner("開啟 Excel 合併中，請稍候…"):
                    try:
                        merged_bytes = merge_sheet_into_bytes(photos, project_name, preset_key, target_file.getvalue())
                        st.success("合併完成，請下載。")
                        st.download_button(
                            "下載合併後的結算明細表",
                            data=merged_bytes,
                            file_name=target_file.name.rsplit(".", 1)[0] + "_已合併.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            key=f"{ns}_merge_download",
                        )
                    except Exception as e:
                        st.error(f"合併失敗：{e}")
    else:
        st.info("請先在上面的格子裡上傳照片。")
