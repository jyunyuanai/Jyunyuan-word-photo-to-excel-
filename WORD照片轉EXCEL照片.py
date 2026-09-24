# -*- coding: utf-8 -*-
"""
Word 施工照片 → 工程隱蔽照片 Excel
流程：上傳多個 .docx → 解析日期/檢查項目/設計值/實測值/照片 → 勾選要輸出的照片 → 產出 Excel

安裝：pip install streamlit python-docx openpyxl pillow
執行：streamlit run WORD照片轉EXCEL照片.py
"""
import io
import os
import re
import shutil
import tempfile
import hashlib

import streamlit as st
from _auth import require_password
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

LABELS = {"日期": "date", "檢查項目": "item", "設計值": "design", "實測值": "measured"}


# ---------------------------------------------------------------- 解析 Word
VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _images_in_element(element, part):
    """回傳 element 內所有圖片的 bytes（依出現順序），
    同時支援新版 DrawingML（a:blip）與舊版 VML（v:imagedata）兩種內嵌方式"""
    blobs = []
    for blip in element.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid and rid in part.related_parts:
            blobs.append(part.related_parts[rid].blob)
    for imgdata in element.findall(".//" + VML_IMAGEDATA):
        rid = imgdata.get(R_ID)
        if rid and rid in part.related_parts:
            blobs.append(part.related_parts[rid].blob)
    return blobs


def _cell_text(cell):
    return "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def _split_label(text):
    """'日期：115年7月18日' → ('日期', '115年7月18日')"""
    m = re.match(r"\s*([^\s:：]+)[:：]\s*(.*)", text)
    if m:
        return m.group(1), m.group(2).strip()
    return None, text


def _date_sort_key(date_text):
    """把『115年7月6日』『115年06月25日』等民國日期轉成可排序的 tuple；
    格式不明或空白排到最後"""
    m = re.search(r"(\d+)\s*年\s*(\d+)\s*月\s*(\d+)\s*日", date_text or "")
    if m:
        return (0, int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return (1, 0, 0, 0)


def _unique_cells(row):
    """跳過因合併儲存格而重複出現的 cell"""
    cells = row.cells
    out = []
    for c, cell in enumerate(cells):
        if c > 0 and cell._tc is cells[c - 1]._tc:
            continue
        out.append(cell)
    return out


def _to_png(blob):
    """轉成 PNG bytes；EMF/WMF 等 PIL 無法讀的格式回傳 None"""
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
        im = ImageOps.exif_transpose(im)
        if im.height > im.width:
            im = im.rotate(-90, expand=True)  # 一律輸出橫式
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGB")
        out = io.BytesIO()
        im.save(out, format="PNG")
        return out.getvalue(), im.size
    except Exception:
        return None, None


def _parse_structured_table(table, doc, filename, photos):
    """解析『日期／檢查項目／設計值＋實測值／照片』每 4 列一組的表格。
    找不到符合格式的起始列則回傳 False，交給呼叫端做一般解析。"""
    rows = table.rows
    n = len(rows)
    matched_any = False
    r = 0
    while r < n:
        first_cells = _unique_cells(rows[r])
        first_text = first_cells[0].text.strip() if first_cells else ""
        label, _ = _split_label(first_text)
        if label != "日期" or r + 3 >= n:
            r += 1
            continue

        fields = {"date": "", "item": "", "design": "", "measured": ""}
        for cell in first_cells:
            lb, val = _split_label(cell.text.strip())
            if lb in LABELS:
                fields[LABELS[lb]] = val

        for offset in (1, 2):
            for cell in _unique_cells(rows[r + offset]):
                lb, val = _split_label(cell.text.strip())
                if lb in LABELS:
                    fields[LABELS[lb]] = val

        photo_row = rows[r + 3]
        blobs = []
        for cell in _unique_cells(photo_row):
            blobs.extend(_images_in_element(cell._tc, doc.part))
        if not blobs:
            # 有些版面照片緊接在下一列
            if r + 4 < n:
                for cell in _unique_cells(rows[r + 4]):
                    blobs.extend(_images_in_element(cell._tc, doc.part))

        if blobs:
            matched_any = True
            for blob in blobs:
                photos.append({"blob": blob, **fields})

        r += 4
    return matched_any


def parse_docx(file_bytes, filename):
    """依文件閱讀順序取出照片；優先辨識『日期/檢查項目/設計值/實測值』結構化表格，
    辨識不出來時退回一般的圖片＋說明啟發式解析。"""
    doc = Document(io.BytesIO(file_bytes))
    part = doc.part
    body = doc.element.body

    project_name, site = "", ""
    photos = []

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, doc).text.strip()
            label, val = _split_label(text)
            if label == "工程名稱":
                project_name = val
            elif label == "施工地點":
                site = val
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            if _parse_structured_table(table, doc, filename, photos):
                continue

            # 一般解析：表格內圖片＋鄰近儲存格文字當說明
            rows = table.rows
            for r, row in enumerate(rows):
                cells = row.cells
                for c, cell in enumerate(cells):
                    if c > 0 and cell._tc is cells[c - 1]._tc:
                        continue
                    blobs = _images_in_element(cell._tc, part)
                    if not blobs:
                        continue
                    caption = _cell_text(cell)
                    if not caption and r + 1 < len(rows):
                        below = rows[r + 1].cells
                        if c < len(below):
                            caption = _cell_text(below[c])
                    if not caption and c + 1 < len(cells):
                        caption = _cell_text(cells[c + 1])
                    for blob in blobs:
                        photos.append({"blob": blob, "item": caption, "date": "", "design": "", "measured": ""})

    # 轉 PNG、編號
    result = []
    for i, p in enumerate(photos, 1):
        png, size = _to_png(p["blob"])
        if png is None:
            continue
        result.append({
            "id": f"{filename}#{i}",
            "source": filename,
            "index": i,
            "png": png,
            "size": size,
            "date": p.get("date", ""),
            "item": p.get("item", ""),
            "design": p.get("design", ""),
            "measured": p.get("measured", ""),
            "selected": False,
        })
    return result, project_name, site


def parse_xlsx(file_bytes, filename):
    """讀取本工具（或格式相同）產出的『日期/檢查項目/設計值/實測值+照片』Excel，
    把照片與對應欄位還原成跟 parse_docx 相同結構的資料，方便繼續編輯/合併。"""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = None
    for name in ("隱蔽照片", "工程抽查照片"):
        if name in wb.sheetnames:
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    def cell_text(r, c):
        v = ws.cell(row=r, column=c).value
        return str(v).strip() if v is not None else ""

    project_name, site = "", ""
    for r in range(1, min(10, ws.max_row) + 1):
        label = cell_text(r, 1).rstrip("：:")
        if label == "工程名稱":
            project_name = cell_text(r, 3)
        elif label == "施工地點":
            site = cell_text(r, 3)

    images = [im for im in getattr(ws, "_images", []) if im.anchor is not None]
    images.sort(key=lambda im: im.anchor._from.row)

    result = []
    i = 0
    for img in images:
        photo_row = img.anchor._from.row + 1  # openpyxl anchor row 是 0-index
        date = cell_text(photo_row - 3, 3)
        item = cell_text(photo_row - 2, 3)
        design = cell_text(photo_row - 1, 3)
        measured = cell_text(photo_row - 1, 9)
        try:
            blob = img._data()
        except Exception:
            continue
        png, size = _to_png(blob)
        if png is None:
            continue
        i += 1
        result.append({
            "id": f"{filename}#{i}",
            "source": filename,
            "index": i,
            "png": png,
            "size": size,
            "date": date,
            "item": item,
            "design": design,
            "measured": measured,
            "selected": False,
        })
    return result, project_name, site


# ---------------------------------------------------------------- 輸出 Excel
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LABEL_FONT = Font(name="標楷體", size=14, bold=False, charset=136, family=4)
VALUE_FONT = Font(name="標楷體", size=14, charset=136, family=4)
TITLE_FONT = Font(name="標楷體", size=18, bold=True, charset=136, family=4)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LABEL_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)

NCOLS = 12                  # A~L 為內容欄，M 為照片編號欄
COL_WIDTH = 7.125
IDX_COL_WIDTH = 8.875
ROW_H_HEADER = 24.95
ROW_H_LABEL = 23.1
ROW_H_IMG = 285
IMG_PAD_PX = 8               # 保留一點邊界，避免欄寬/列高換算誤差讓圖片頂出格線


def _excel_col_width_px(chars, mdw=7):
    """依 Excel 官方公式，把「字元寬度」換算成實際像素寬度（預設字型 Calibri 11，MDW=7px）"""
    return int((256 * chars + int(128 / mdw)) / 256 * mdw)


BOX_W_PX = _excel_col_width_px(COL_WIDTH) * NCOLS
BOX_H_PX = int(ROW_H_IMG * 4 / 3)
MAX_W_PX = BOX_W_PX - IMG_PAD_PX
MAX_H_PX = BOX_H_PX - IMG_PAD_PX


def _fit_image(png_bytes):
    im = Image.open(io.BytesIO(png_bytes))
    w, h = im.size
    scale = min(MAX_W_PX / w, MAX_H_PX / h, 1.0)
    im = im.resize((max(1, int(w * scale)), max(1, int(h * scale))))
    out = io.BytesIO()
    im.save(out, format="PNG")
    out.seek(0)
    return out


def _add_centered_image(ws, png_bytes, row, col=1):
    """把圖片縮放後，置中貼在 col 欄、row 列所在的那個（合併）格子裡"""
    img = XLImage(_fit_image(png_bytes))
    off_x = max(0, (BOX_W_PX - img.width) // 2)
    off_y = max(0, (BOX_H_PX - img.height) // 2)
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off_x), row=row - 1, rowOff=pixels_to_EMU(off_y)),
        ext=XDRPositiveSize2D(cx=pixels_to_EMU(img.width), cy=pixels_to_EMU(img.height)),
    )
    ws.add_image(img)


def _label_value_row(ws, row, last, label, value, label_range, value_range, border=True):
    ws[f"{label_range[0]}{row}"] = f"{label}：" if label else ""
    ws[f"{label_range[0]}{row}"].font = LABEL_FONT
    ws[f"{label_range[0]}{row}"].alignment = LABEL_ALIGN
    if label_range[0] != label_range[1]:
        ws.merge_cells(f"{label_range[0]}{row}:{label_range[1]}{row}")
    ws[f"{value_range[0]}{row}"] = value
    ws[f"{value_range[0]}{row}"].font = VALUE_FONT
    ws[f"{value_range[0]}{row}"].alignment = LEFT
    if value_range[0] != value_range[1]:
        ws.merge_cells(f"{value_range[0]}{row}:{value_range[1]}{row}")
    if border:
        for c in range(1, NCOLS + 1):
            ws.cell(row=row, column=c).border = BORDER


def build_excel(photos, project_name, site):
    wb = Workbook()
    # 範本的預設字型是新細明體 12，Excel 的欄寬字元數是依「預設字型」換算成像素，
    # 字型不同會讓同一個寬度數字顯示出不同的實際寬度，這裡對齊範本的預設字型
    wb._named_styles[0].font = Font(name="新細明體", size=12)
    ws = wb.active
    ws.title = "隱蔽照片"
    for c in range(1, NCOLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTH
    ws.column_dimensions[get_column_letter(NCOLS + 1)].width = IDX_COL_WIDTH

    last = get_column_letter(NCOLS)
    row = 1

    # 頁首：標題／工程名稱／施工地點
    ws.merge_cells(f"A{row}:{last}{row}")
    ws[f"A{row}"] = "工程抽查照片"
    ws[f"A{row}"].font = TITLE_FONT
    ws[f"A{row}"].alignment = CENTER
    ws.row_dimensions[row].height = ROW_H_HEADER
    row += 1

    _label_value_row(ws, row, last, "工程名稱", project_name, ("A", "B"), ("C", last), border=False)
    ws.row_dimensions[row].height = ROW_H_HEADER
    row += 1

    _label_value_row(ws, row, last, "施工地點", site, ("A", "B"), ("C", last), border=False)
    ws.row_dimensions[row].height = ROW_H_HEADER
    row += 1

    for n, p in enumerate(photos, 1):
        ws.row_dimensions[row].height = ROW_H_LABEL
        _label_value_row(ws, row, last, "日期", p.get("date", ""), ("A", "B"), ("C", last))
        row += 1

        ws.row_dimensions[row].height = ROW_H_LABEL
        _label_value_row(ws, row, last, "檢查項目", p.get("item", ""), ("A", "B"), ("C", last))
        row += 1

        ws.row_dimensions[row].height = ROW_H_LABEL
        ws[f"A{row}"] = "設計值："
        ws[f"A{row}"].font = LABEL_FONT
        ws[f"A{row}"].alignment = LABEL_ALIGN
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"C{row}"] = p.get("design", "")
        ws[f"C{row}"].font = VALUE_FONT
        ws[f"C{row}"].alignment = LEFT
        ws.merge_cells(f"C{row}:F{row}")
        ws[f"G{row}"] = "實測值："
        ws[f"G{row}"].font = LABEL_FONT
        ws[f"G{row}"].alignment = LABEL_ALIGN
        ws.merge_cells(f"G{row}:H{row}")
        ws[f"I{row}"] = p.get("measured", "")
        ws[f"I{row}"].font = VALUE_FONT
        ws[f"I{row}"].alignment = LEFT
        ws.merge_cells(f"I{row}:{last}{row}")
        for c in range(1, NCOLS + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

        ws.row_dimensions[row].height = ROW_H_IMG
        ws.merge_cells(f"A{row}:{last}{row}")
        for c in range(1, NCOLS + 1):
            ws.cell(row=row, column=c).border = BORDER
        _add_centered_image(ws, p["png"], row)
        ws.cell(row=row, column=NCOLS + 1, value=n).alignment = CENTER
        row += 1

    ws.print_area = f"A1:{get_column_letter(NCOLS + 1)}{max(row - 1, 1)}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.59
    ws.page_margins.top = 0.79
    ws.page_margins.bottom = 0.59
    ws.print_options.horizontalCentered = True

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# ---------------------------------------------------------------- 合併進既有結算明細表
def merge_sheet_into_bytes(photos, project_name, site, target_bytes, sheet_name="隱蔽照片"):
    """用 Excel COM 把新產生的分頁取代掉 target_bytes（既有結算明細表）裡的同名分頁；
    其餘分頁與分頁順序不動。回傳合併後的檔案 bytes，原始上傳內容完全不會被修改。"""
    try:
        import pythoncom
        import win32com.client as win32
    except ImportError:
        raise RuntimeError(
            "載入 pywin32 失敗。請先安裝（uv pip install pywin32 或 pip install pywin32），"
            "安裝後必須「重新啟動」這個網頁程式才會生效。"
        )

    tmp_dir = tempfile.mkdtemp(prefix="word2excel_")
    try:
        src_path = os.path.join(tmp_dir, "source.xlsx")
        with open(src_path, "wb") as f:
            f.write(build_excel(photos, project_name, site).read())

        tgt_path = os.path.join(tmp_dir, "target.xlsx")
        with open(tgt_path, "wb") as f:
            f.write(target_bytes)

        pythoncom.CoInitialize()
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        try:
            src_wb = excel.Workbooks.Open(src_path)
            tgt_wb = excel.Workbooks.Open(tgt_path)
            try:
                idx = None
                for i in range(1, tgt_wb.Sheets.Count + 1):
                    if tgt_wb.Sheets(i).Name == sheet_name:
                        idx = i
                        break
                if idx is not None:
                    tgt_wb.Sheets(idx).Delete()

                src_sheet = src_wb.Sheets(1)
                if idx is not None and idx <= tgt_wb.Sheets.Count:
                    src_sheet.Copy(Before=tgt_wb.Sheets(idx))
                    new_sheet = tgt_wb.Sheets(idx)
                else:
                    src_sheet.Copy(After=tgt_wb.Sheets(tgt_wb.Sheets.Count))
                    new_sheet = tgt_wb.Sheets(tgt_wb.Sheets.Count)
                new_sheet.Name = sheet_name

                tgt_wb.Save()
            finally:
                tgt_wb.Close(False)
                src_wb.Close(False)
        finally:
            excel.Quit()
            pythoncom.CoUninitialize()

        with open(tgt_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ---------------------------------------------------------------- Streamlit UI
st.set_page_config(page_title="Word 照片 → 工程隱蔽照片 Excel", layout="wide")
require_password()
st.title("Word 施工照片 → 工程隱蔽照片 Excel")

with st.sidebar:
    st.header("設定")
    cols_per_row = st.slider("縮圖每列張數", 2, 8, 8)

u1, u2 = st.columns(2)
with u1:
    uploaded = st.file_uploader(
        "① 上傳照片來源檔（Word .docx，或本工具產出格式的 .xlsx，可多選、可混合）",
        type=["docx", "xlsx"], accept_multiple_files=True)
with u2:
    target_file = st.file_uploader("② 上傳結算明細表.xlsx（要把照片合併進去的檔案）",
                                   type=["xlsx"], key="target_xlsx")

if uploaded:
    # 檔案組合改變時重新解析
    sig = hashlib.md5("".join(f.name + str(f.size) for f in uploaded).encode()).hexdigest()
    if st.session_state.get("sig") != sig:
        photos = []
        project_name, site = "", ""
        with st.spinner("解析檔案中…"):
            for f in uploaded:
                try:
                    parser = parse_xlsx if f.name.lower().endswith(".xlsx") else parse_docx
                    got, pn, st_ = parser(f.read(), f.name)
                    photos.extend(got)
                    project_name = project_name or pn
                    site = site or st_
                except Exception as e:
                    st.error(f"{f.name} 解析失敗：{e}")
        for p in photos:
            p["selected"] = False
        st.session_state.photos = photos
        st.session_state.sig = sig
        st.session_state.project_name = project_name
        st.session_state.site = site

    photos = st.session_state.photos

    with st.sidebar:
        project_name = st.text_input("工程名稱", st.session_state.get("project_name", ""))
        site = st.text_input("施工地點", st.session_state.get("site", ""))

    st.caption(f"共解析 {len(photos)} 張照片，來自 {len(uploaded)} 個檔案")

    # 縮圖格（依日期排序顯示）
    photos = sorted(photos, key=lambda p: _date_sort_key(p["date"]))
    for start in range(0, len(photos), cols_per_row):
        cols = st.columns(cols_per_row)
        for col, p in zip(cols, photos[start:start + cols_per_row]):
            with col:
                st.image(p["png"], use_container_width=True)
                st.caption(f"{p['source']}　#{p['index']}")
                p["selected"] = st.checkbox("納入輸出", value=p["selected"], key=f"sel_{p['id']}")
                p["date"] = st.text_input("日期", p["date"], key=f"date_{p['id']}")
                p["item"] = st.text_area("檢查項目", p["item"], key=f"item_{p['id']}", height=70)
                d1, d2 = st.columns(2)
                p["design"] = d1.text_input("設計值", p["design"], key=f"design_{p['id']}")
                p["measured"] = d2.text_input("實測值", p["measured"], key=f"measured_{p['id']}")

    selected = [p for p in photos if p["selected"]]
    selected.sort(key=lambda p: _date_sort_key(p["date"]))
    st.divider()
    st.subheader(f"已勾選 {len(selected)} 張照片（已依日期排序）")
    if selected:
        xlsx = build_excel(selected, project_name, site)
        st.download_button(
            "產生並下載 Excel",
            data=xlsx,
            file_name=(project_name or "工程抽查照片") + ".xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        st.divider()
        st.subheader("合併到既有結算明細表")
        if target_file is None:
            st.info("要合併的話，請在最上方 ② 上傳結算明細表.xlsx。")
        else:
            st.caption(f"會用目前勾選的 {len(selected)} 張照片，取代掉「{target_file.name}」裡「隱蔽照片」那個分頁，"
                       "其他分頁不動，完成後另外下載，不會動到你上傳的原始檔案。需要這台電腦有安裝 Microsoft Excel。")
            if st.button("合併並產生新檔案", type="primary"):
                with st.spinner("開啟 Excel 合併中，請稍候…"):
                    try:
                        merged_bytes = merge_sheet_into_bytes(selected, project_name, site, target_file.getvalue())
                        st.success("合併完成，請下載。")
                        st.download_button(
                            "下載合併後的結算明細表",
                            data=merged_bytes,
                            file_name=target_file.name.rsplit(".", 1)[0] + "_已合併隱蔽照片.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    except Exception as e:
                        st.error(f"合併失敗：{e}")
else:
    st.info("請先上傳一個或多個 .docx 或 .xlsx 照片檔案。")
